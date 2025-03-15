# Importing Used Libraries
import random
import timeit
import openpyxl
import pandas
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.workbook import Workbook
import matplotlib.pyplot as plt

# Importing Sorting Algorithms
from Sorting_Implementation.Bubble_Sort import Bubble_Sort as BS
from Sorting_Implementation.Insertion_Sort import Insertion_Sort as IS
from Sorting_Implementation.Selection_Sort import Selection_Sort as SS
from Sorting_Implementation.Quicksort import timedQuickSort as QS
from Sorting_Implementation.Quicksort import kthElement as QSK
from Sorting_Implementation.Hybrid_Merge_Sort import merge_sort as HMS
from Sorting_Implementation.Merge_Sort import timedMergeSort as MS
from Sorting_Implementation.Heap_Sort import heapsort as HS


# Generating Random Array Based On Needed Length


def random_array(length):
    arr = []
    for i in range(0, int(length)):
        arr.insert(0, random.randint(0, 100))
    return arr


# Validating Sorting Algorithms

arr = random_array(10)

print("\nInitial Array =>", arr)
result = BS(arr[:])
print("Bubble Sort =>", result[0])
print("Time Taken By Bubble Sort: ", result[1])
result = IS(arr[:])
print("Insertion Sort =>", result[0])
print("Time Taken By Insertion Sort:  ", result[1])
result = SS(arr[:])
print("Selection Sort =>", result[0])
print("Time Taken By Selection Sort:  ", result[1])
result = QS(arr[:])
print("Quick Sort =>", result[0])
print("Time Taken By Quick Sort:  ", result[1])
result = MS(arr[:])
print("Merge Sort =>", result[0])
print("Time Taken By Merge Sort:  ", result[1])
result = HS(arr[:])
print("Heap Sort =>", result[0])
print("Time Taken By Heap Sort:  ", result[1])

# print("Kth Element => Input Array [3,41,16,25,63,52,40]. 3rd smallest element:")
# array = [3,41,16,25,63,52,40]
# print(QSK(array,3, 0, len(array) - 1))
# print("Kth Element => Input Array [23,78,45,8,32,56]. 2nd smallest element:")
# array = [23,78,45,8,32,56]
# print(QSK(array,2, 0, len(array) - 1))
print("Hybrid Merge Sort => Input Array [8, 6, 4, 8, 3, 6, 1]. Threshold = 5:")
array = [8,6,4,8,3,6,1]
print(HMS(array, 5))
# Displaying Time Taken Based On Each Input Size


def time_complexity(input_sizes):
    time_bubble = []
    time_insertion = []
    time_selection = []
    for i in range(0, len(input_sizes)):
        arr = random_array(input_sizes[i])
        time_bubble.append(BS(arr)[1])
        time_insertion.append(IS(arr)[1])
        time_selection.append(SS(arr)[1])
    return time_bubble, time_insertion, time_selection


input_sizes = [0, 2500, 5000, 7500, 10000]
time_ = time_complexity(input_sizes)
bubble_sort_time = time_[0]
insertion_sort_time = time_[1]
selection_sort_time = time_[2]

# Ploting Graphs Using matplotlib
plt.plot(input_sizes, bubble_sort_time, marker='o',
         linestyle='-', color='r', label='Bubble Sort')
plt.plot(input_sizes, insertion_sort_time, marker='o',
         linestyle='-', color='g', label='Insertion Sort')
plt.plot(input_sizes, selection_sort_time, marker='o',
         linestyle='-', color='b', label='Selection Sort')

plt.xlabel("Input Size(N)")
plt.ylabel("Execution Time (microsec)")
plt.title("Time Complexity Graph")
plt.legend()
plt.show()

# Ploting Graphs On Excel Using pandas And openpyxl


def entireAlgorithm(runs):
    resultBubble = [0] * runs
    resultInsertion = [0] * runs
    resultSelection = [0] * runs
    resultQuick = [0] * runs
    resultMerge = [0] * runs
    resultHeap = [0] * runs
    elementNumbers = [0] * runs
    for i in range(runs):
        arrayElements = random.randint(1, 1000)
        if arrayElements in elementNumbers:
            i -= 1
        else:
            elementNumbers[i] = arrayElements
            array = random_array(arrayElements)
            runResult = BS(array[:])
            resultBubble[i] = runResult[1]
            runResult = IS(array[:])
            resultInsertion[i] = runResult[1]
            runResult = SS(array[:])
            resultSelection[i] = runResult[1]
            runResult = QS(array[:])
            resultQuick[i] = runResult[1]
            runResult = MS(array[:])
            resultMerge[i] = runResult[1]
            runResult = HS(array[:])
            resultHeap[i] = runResult[1]
    return elementNumbers, resultBubble, resultInsertion, resultSelection, resultQuick, resultMerge, resultHeap


runResult = entireAlgorithm(500)
algorithRunTimeData = pandas.DataFrame(
    {'Input Size (n)': runResult[0], 'Bubble Sort': runResult[1],  'Insertion Sort': runResult[2], 'Selection Sort': runResult[3], 'Quick Sort': runResult[4], 'Merge Sort': runResult[5], 'Heap Sort': runResult[6]})
filename = 'AlgorithmRuntime.xlsx'
algorithRunTimeData.to_excel(filename)
wb = Workbook()
ws = wb.active
for rowIndex, row in enumerate(algorithRunTimeData.values, start=2):
    for columnIndex, value in enumerate(row, start=1):
        ws.cell(row=rowIndex, column=columnIndex, value=value)
chart = ScatterChart()
chart.x_axis.title = 'Input Size (n)'
chart.y_axis.title = 'Execution Time (microseconds)'
chart.scatterStyle = 'marker'
x_axis = Reference(ws, min_col=1, max_col=1, min_row=2, max_row=501)
for i, name in enumerate(["Bubble Sort", "Insertion Sort", "Selection Sort", "Quick Sort", "Merge Sort", "Heap Sort"], start=2):
    y_axis = Reference(ws, min_col=i, min_row=2, max_row=501)
    series = Series(y_axis, x_axis, title=name)
    series.marker.symbol = "circle"
    series.graphicalProperties.line.noFill = True
    chart.series.append(series)

ws.add_chart(chart, "J8")
wb.save(filename)
