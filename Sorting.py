import random
import timeit
import openpyxl
import pandas as pd
from openpyxl.chart import ScatterChart, Series,Reference
from openpyxl.workbook import Workbook


def bubbleSort(array):
    start = timeit.default_timer()
    for i in range(len(array)):
        changed = False
        for j in range(len(array) - i - 1):
            if array[j] > array[j + 1]:
                array[j], array[j+1] =array[j+1],array[j]
                changed = True
        if not changed:
            break
    end = timeit.default_timer()
    time = 1000000*(end - start)
    return array, time

def insertionSort(array):
    start = timeit.default_timer()
    for i in range(1, len(array)-1):
        j = i
        while j>0 and array[j-1]>array[j]:
            array[j], array[j-1] = array[j-1], array[j]
            j -= 1
    end = timeit.default_timer()
    time = 1000000*(end - start)
    return array, time

def selectionSort(array):
    start = timeit.default_timer()
    for i in range(len(array)-1):
        minimum = i
        for j in range(i+1, len(array)):
            if array[j]<array[minimum]:
                minimum =j
            if minimum != i:
                array[i], array[minimum] = array[minimum], array[i]
    end = timeit.default_timer()
    time = 1000000*(end - start)
    return array, time

def generateArray(size):
    array = [0]*size
    for i in range(size):
        array[i] = random.randint(0, 200)
    return array

arr = generateArray(100)

print("Initial Array:", arr)
result = bubbleSort(arr)
print("Bubble Sort: ", result[0])
print("Time Taken By Bubble Sort: ", result[1])
result = insertionSort(arr)
print("Insertion Sort: ", result[0])
print("Time Taken By Insertion Sort: ", result[1])
result = selectionSort(arr)
print("Selection Sort: ", result[0])
print("Time Taken By Selection Sort: ", result[1])

def entireAlgorithm(runs):
    resultBubble = [0] * runs
    resultInsertion = [0] * runs
    resultSelection = [0] * runs
    elementNumbers = [0] * runs
    for i in range(runs):
        arrayElements = random.randint(1, 1000)
        if arrayElements in elementNumbers:
            i -= 1;
        else:
            elementNumbers[i] = arrayElements
            array = generateArray(arrayElements)
            runResult = bubbleSort(array)
            resultBubble[i] = runResult[1]
            runResult = insertionSort(array)
            resultInsertion[i] = runResult[1]
            runResult = selectionSort(array)
            resultSelection[i] = runResult[1]
    return elementNumbers, resultBubble, resultInsertion, resultSelection

runResult = entireAlgorithm(500)
algorithRunTimeData = pd.DataFrame({'Input Size (n)' : runResult[0], 'Bubble Sort' : runResult[1],  'Insertion Sort' : runResult[2], 'Selection Sort' : runResult[3]})
filename = 'AlgorithmRuntime.xlsx'
algorithRunTimeData.to_excel(filename)
wb = Workbook()
ws = wb.active
for rowIndex, row in enumerate(algorithRunTimeData.values, start=2):
    for columnIndex, value in enumerate(row, start=1):
        ws.cell(row=rowIndex, column=columnIndex,value = value)
chart = ScatterChart()
chart.x_axis.title = 'Input Size (n)'
chart.y_axis.title = 'Excution Time (microseconds)'
chart.scatterStyle = 'marker'
x_axis = Reference(ws, min_col=1, max_col=1, min_row=1, max_row=501)
for i, name in enumerate(["Bubble Sort", "Insertion Sort", "Selection Sort"], start = 2):
    y_axis = Reference(ws, min_col=i, min_row=2, max_row=501)
    series = Series(y_axis, x_axis, title = name)
    series.marker.symbol = "circle"
    series.graphicalProperties.line.noFill = True
    chart.series.append(series)

ws.add_chart(chart, "H8")
wb.save(filename)
